# -*- coding: UTF-8 -*-
"""给 Excel 每一列做一次「体检」，输出列级统计画像。

适用场景：拿到一份别人整理的表，先想知道它靠不靠谱——哪列缺失严重、哪列其实是文本混了数字、
数值列的量级和离散程度如何、有没有明显的异常值。手工在 Excel 里拖公式既慢又容易看漏。

与 office.api.excel.find_excel_data 的区别：find_excel_data 是「按关键词找内容」，
本方法是「按列做整体统计」，用于数据质量初筛。

Author:
    my12yyds

Project:
    https://www.python-office.com
"""

import datetime
import os
import statistics

import openpyxl

#: 判定为数值类型的 Python 类型。
_NUMBER_TYPES = (int, float)
#: 判定为日期类型的 Python 类型。
_DATE_TYPES = (datetime.datetime, datetime.date, datetime.time)


def excel_column_stats(excel_path: str, sheet_name: str = None, top_n: int = 5) -> dict:
    """统计每一列的填充率、类型、唯一值数量与数值分布。

    Args:
        excel_path (str): 需要统计的 Excel 文件路径，支持 .xlsx。
        sheet_name (str, optional): 工作表名称。默认 None，表示第一个工作表。
        top_n (int, optional): 文本列返回出现次数最多的前 N 个值，设为 0 表示不返回。
            默认 5。

    Returns:
        dict: 统计结果，包含以下键：
            - sheet (str): 工作表名称。
            - row_count (int): 数据行数（不含表头）。
            - column_count (int): 列数。
            - fill_rate (float): 整体填充率，非空单元格占比，保留 4 位小数。
            - columns (list): 每一列的统计，按列顺序排列。每项为 dict，含
              name（列名）、index（列序号，从 1 开始）、filled（非空数）、empty（空值数）、
              empty_ratio（空值占比）、unique（去重后的非空值个数）、type（推断类型，
              取值 number / text / date / mixed / empty），并按类型附带以下之一：
                - numeric (dict): 数值列的 min、max、mean、median、stdev，
                  样本量不足 2 时 stdev 为 None。
                - date_range (dict): 日期列的 min、max，为 ISO 格式字符串。
                - top (list): 文本列出现最多的值，每项含 value 和 count。

    Raises:
        FileNotFoundError: Excel 文件不存在。
        ValueError: 工作表不存在、表格为空，或 top_n 为负数。
    """
    if not os.path.isfile(excel_path):
        raise FileNotFoundError('文件不存在：{}'.format(excel_path))
    if top_n < 0:
        raise ValueError('top_n 不能为负数：{}'.format(top_n))

    header, rows = _load_sheet(excel_path, sheet_name)
    if not header:
        raise ValueError('表格为空，至少需要一行表头：{}'.format(excel_path))

    filled_total = 0
    cell_total = 0
    columns = []

    for index, name in enumerate(header):
        values = []
        for row in rows:
            value = row[index] if index < len(row) else None
            cell_total += 1
            if _is_blank(value):
                continue
            filled_total += 1
            values.append(value)

        empty = len(rows) - len(values)
        column = {
            'name': name,
            'index': index + 1,
            'filled': len(values),
            'empty': empty,
            'empty_ratio': round(empty / len(rows), 4) if rows else 0.0,
            'unique': len({_hashable(value) for value in values}),
            'type': _infer_type(values),
        }

        if column['type'] == 'number':
            column['numeric'] = _numeric_stats(values)
        elif column['type'] == 'date':
            column['date_range'] = _date_range(values)
        elif column['type'] in ('text', 'mixed') and top_n:
            column['top'] = _top_values(values, top_n)

        columns.append(column)

    return {
        'sheet': _sheet_label(excel_path, sheet_name),
        'row_count': len(rows),
        'column_count': len(header),
        'fill_rate': round(filled_total / cell_total, 4) if cell_total else 0.0,
        'columns': columns,
    }


def _load_sheet(path, sheet_name):
    """读取工作表，返回（表头列表，数据行列表）。

    Args:
        path (str): Excel 文件路径。
        sheet_name (str): 工作表名称，None 表示第一个工作表。

    Returns:
        tuple: (header, rows)。rows 已去掉末尾的空行。

    Raises:
        ValueError: 指定的工作表不存在。
    """
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        if sheet_name is None:
            worksheet = workbook.worksheets[0]
        else:
            if sheet_name not in workbook.sheetnames:
                raise ValueError('工作表不存在：{}，可选：{}'.format(sheet_name, workbook.sheetnames))
            worksheet = workbook[sheet_name]
        raw_rows = [list(row) for row in worksheet.iter_rows(values_only=True)]
    finally:
        workbook.close()

    while raw_rows and all(cell is None or cell == '' for cell in raw_rows[-1]):
        raw_rows.pop()
    if not raw_rows:
        return [], []

    header = []
    for index, cell in enumerate(raw_rows[0]):
        text = str(cell).strip() if cell is not None else ''
        header.append(text or '列{}'.format(index + 1))
    return header, raw_rows[1:]


def _sheet_label(path, sheet_name):
    """返回用于结果展示的工作表名称。

    Args:
        path (str): Excel 文件路径。
        sheet_name (str): 用户指定的工作表名称。

    Returns:
        str: 工作表名称。
    """
    if sheet_name:
        return sheet_name
    workbook = openpyxl.load_workbook(path, read_only=True)
    try:
        return workbook.sheetnames[0]
    finally:
        workbook.close()


def _is_blank(value):
    """判断单元格是否为空。

    Args:
        value (object): 单元格值。

    Returns:
        bool: None 或去空格后的空字符串视为空。
    """
    return value is None or (isinstance(value, str) and not value.strip())


def _hashable(value):
    """把单元格值转成可用于去重计数的形式。

    Args:
        value (object): 单元格值。

    Returns:
        object: 字符串去除首尾空格；其余类型原样返回。
    """
    if isinstance(value, str):
        return value.strip()
    return value


def _infer_type(values):
    """推断一列的类型。

    Args:
        values (list): 该列的非空值列表。

    Returns:
        str: number、text、date、mixed 或 empty。
    """
    if not values:
        return 'empty'
    flags = set()
    for value in values:
        if isinstance(value, bool):
            flags.add('text')
        elif isinstance(value, _NUMBER_TYPES):
            flags.add('number')
        elif isinstance(value, _DATE_TYPES):
            flags.add('date')
        elif isinstance(value, str):
            flags.add('text')
        else:
            flags.add('mixed')
    if len(flags) == 1:
        return flags.pop()
    return 'mixed'


def _numeric_stats(values):
    """计算数值列的分布指标。

    Args:
        values (list): 该列的数值列表。

    Returns:
        dict: 含 min、max、mean、median、stdev；stdev 在样本量小于 2 时为 None。
    """
    numbers = [float(value) for value in values if isinstance(value, _NUMBER_TYPES) and not isinstance(value, bool)]
    if not numbers:
        return {'min': None, 'max': None, 'mean': None, 'median': None, 'stdev': None}
    return {
        'min': round(min(numbers), 4),
        'max': round(max(numbers), 4),
        'mean': round(statistics.mean(numbers), 4),
        'median': round(statistics.median(numbers), 4),
        'stdev': round(statistics.stdev(numbers), 4) if len(numbers) > 1 else None,
    }


def _date_range(values):
    """计算日期列的最早与最晚值。

    Args:
        values (list): 该列的日期值列表。

    Returns:
        dict: 含 min 和 max 两个 ISO 格式字符串。
    """
    dates = [value for value in values if isinstance(value, _DATE_TYPES)]
    if not dates:
        return {'min': None, 'max': None}
    return {
        'min': min(dates).isoformat(),
        'max': max(dates).isoformat(),
    }


def _top_values(values, top_n):
    """统计出现次数最多的值。

    Args:
        values (list): 该列的非空值列表。
        top_n (int): 返回的条数。

    Returns:
        list: 每项含 value 和 count，按次数降序、同次数按值升序排列。
    """
    counts = {}
    for value in values:
        key = _hashable(value)
        counts[key] = counts.get(key, 0) + 1
    ordered = sorted(counts.items(), key=lambda item: (-item[1], str(item[0])))
    return [{'value': value, 'count': count} for value, count in ordered[:top_n]]
