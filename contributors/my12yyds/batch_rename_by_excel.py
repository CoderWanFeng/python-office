# -*- coding: UTF-8 -*-
"""按 Excel 表格批量重命名文件。

适用场景：手里有一批需要改名的文件，改名前后的对应关系已经整理在 Excel 里。
只需要准备两列（原文件名、新文件名），程序按这个映射批量重命名指定目录下的文件。

与 office.api.file.replace4filename 的区别：replace4filename 是「按规则批量替换
文件名中的字符」，本方法是「按 Excel 中逐行指定的映射改名」，两者互补。

Author:
    my12yyds

Project:
    https://www.python-office.com
"""

import os

import openpyxl


def batch_rename_by_excel(
    excel_path: str,
    folder_path: str,
    old_name_column: str = '原文件名',
    new_name_column: str = '新文件名',
    sheet_name: str = None,
    dry_run: bool = False,
) -> dict:
    """按 Excel 中记录的新旧文件名映射，批量重命名目录下的文件。

    Args:
        excel_path (str): 记录新旧文件名映射的 Excel 文件路径，支持 .xlsx。
        folder_path (str): 需要重命名文件所在的目录路径，只处理该目录第一层的文件。
        old_name_column (str, optional): Excel 中存放原文件名的列名。默认 '原文件名'。
        new_name_column (str, optional): Excel 中存放新文件名的列名。默认 '新文件名'。
        sheet_name (str, optional): 工作表名称。默认 None，表示使用第一个工作表。
        dry_run (bool, optional): 为 True 时只返回将要执行的重命名计划，不真正修改文件。
            默认 False。

    Returns:
        dict: 重命名结果，包含以下 4 个键：
            - renamed (list): 已重命名（或 dry_run 下计划重命名）的记录，
              元素为 (原文件名, 新文件名) 元组。
            - skipped (list): 未执行重命名的记录，元素为 (原文件名, 原因) 元组。
            - missing (list): Excel 中有记录、但目录下找不到的原文件名。
            - conflicts (list): 目标文件名已存在，为避免覆盖而跳过的原文件名。

    Raises:
        FileNotFoundError: Excel 文件不存在，或目录不存在。
        ValueError: Excel 中缺少指定的列名，或工作表不存在。
    """
    if not os.path.isfile(excel_path):
        raise FileNotFoundError('Excel 文件不存在：{}'.format(excel_path))
    if not os.path.isdir(folder_path):
        raise FileNotFoundError('目录不存在：{}'.format(folder_path))

    workbook = openpyxl.load_workbook(excel_path, read_only=True, data_only=True)
    try:
        if sheet_name is None:
            worksheet = workbook.active
        elif sheet_name in workbook.sheetnames:
            worksheet = workbook[sheet_name]
        else:
            raise ValueError('工作表不存在：{}'.format(sheet_name))

        rows = worksheet.iter_rows(values_only=True)
        header = next(rows, None)
        if header is None:
            raise ValueError('Excel 是空的，没有表头行')

        # 表头转成 {列名: 列序号}，用于定位「原文件名」「新文件名」两列
        header_index = {}
        for index, title in enumerate(header):
            if title is not None:
                header_index[str(title).strip()] = index
        for column in (old_name_column, new_name_column):
            if column not in header_index:
                raise ValueError(
                    'Excel 中找不到列「{}」，当前表头是：{}'.format(
                        column, list(header_index.keys())
                    )
                )

        old_index = header_index[old_name_column]
        new_index = header_index[new_name_column]

        renamed = []
        skipped = []
        missing = []
        conflicts = []

        for row in rows:
            old_name = _cell_text(row, old_index)
            new_name = _cell_text(row, new_index)
            # 整行为空（Excel 末尾常见）直接忽略
            if not old_name and not new_name:
                continue
            if not old_name or not new_name:
                skipped.append((old_name or new_name, '原文件名或新文件名为空'))
                continue
            if old_name == new_name:
                skipped.append((old_name, '新旧文件名相同'))
                continue

            source = os.path.join(folder_path, old_name)
            target = os.path.join(folder_path, new_name)
            if not os.path.isfile(source):
                missing.append(old_name)
                continue
            if os.path.exists(target):
                conflicts.append(old_name)
                continue

            if not dry_run:
                os.rename(source, target)
            renamed.append((old_name, new_name))

        return {
            'renamed': renamed,
            'skipped': skipped,
            'missing': missing,
            'conflicts': conflicts,
        }
    finally:
        workbook.close()


def _cell_text(row, index: int) -> str:
    """取出单元格内容并转成去掉首尾空格的字符串。

    Args:
        row (tuple): iter_rows 返回的一行数据。
        index (int): 需要读取的列序号。

    Returns:
        str: 单元格文本，空单元格返回空字符串。
    """
    if index >= len(row):
        return ''
    value = row[index]
    if value is None:
        return ''
    return str(value).strip()


if __name__ == '__main__':
    # 使用前先准备：
    # 1、Excel 中两列表头分别为「原文件名」「新文件名」
    # 2、把要改名的文件放到 folder_path 目录下
    excel_file = './rename_list.xlsx'
    target_folder = './files'

    if not os.path.isfile(excel_file) or not os.path.isdir(target_folder):
        print('演示文件不存在，请先准备：')
        print('  1、在当前目录放一个 {}，两列表头为「原文件名」「新文件名」'.format(excel_file))
        print('  2、在当前目录新建 {} 目录，把要改名的文件放进去'.format(target_folder))
    else:
        result = batch_rename_by_excel(
            excel_path=excel_file,
            folder_path=target_folder,
            dry_run=True,  # 先预览，确认无误后改成 False 真正执行
        )
        print('计划重命名：', result['renamed'])
        print('跳过：', result['skipped'])
        print('找不到的文件：', result['missing'])
        print('目标已存在：', result['conflicts'])
