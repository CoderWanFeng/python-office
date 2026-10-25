# -*- coding: UTF-8 -*-
"""把 Excel 工作表导出为 CSV。

适用场景：把表格交给只会用 CSV 的系统导入、喂给命令行工具，或者需要体积更小的纯文本存档。
中文用户最常踩的坑是编码：直接存 UTF-8 再用 Excel 打开会乱码，所以默认使用
utf-8-sig（带 BOM），双击用 Excel 打开中文正常显示。

与 office.api.excel.excel2pdf 的区别：excel2pdf 输出给人看的 PDF，
本方法输出给程序读的 CSV。

Author:
    my12yyds

Project:
    https://www.python-office.com
"""

import csv
import os
import re

import openpyxl

#: 文件名中需要替换掉的字符。
_INVALID_FILENAME_CHARS = re.compile(r'[\\/:*?"<>|]')


def excel2csv(
    excel_path: str,
    output_dir: str = None,
    sheet_name: str = None,
    encoding: str = 'utf-8-sig',
    delimiter: str = ',',
) -> dict:
    """把一个或多个工作表导出为 CSV 文件。

    Args:
        excel_path (str): Excel 文件路径，支持 .xlsx。
        output_dir (str, optional): CSV 输出目录，目录必须已存在。默认 None，
            表示与 Excel 文件同目录。
        sheet_name (str, optional): 只导出指定的工作表。默认 None，表示导出全部工作表。
        encoding (str, optional): 输出编码。默认 'utf-8-sig'，带 BOM，
            用 Excel 直接打开中文不乱码；需要给 Linux 工具链消费时可改为 'utf-8'。
        delimiter (str, optional): 分隔符。默认 ','；部分地区习惯用 ';' 或制表符 '\\t'。

    Returns:
        dict: 导出结果，包含以下键：
            - files (list): 生成的 CSV 文件绝对路径列表，按工作表顺序排列。
            - rows (dict): 工作表名 -> 写出的数据行数（不含表头）。
            - encoding (str): 实际使用的编码。
            - delimiter (str): 实际使用的分隔符。

    Raises:
        FileNotFoundError: Excel 文件不存在，或输出目录不存在。
        ValueError: 指定的工作表不存在，或 delimiter 为空。
        OSError: 文件写入失败。
    """
    if not os.path.isfile(excel_path):
        raise FileNotFoundError('文件不存在：{}'.format(excel_path))
    if not delimiter:
        raise ValueError('delimiter 不能为空')

    target_dir = output_dir or os.path.dirname(os.path.abspath(excel_path))
    if not os.path.isdir(target_dir):
        raise FileNotFoundError('输出目录不存在：{}'.format(target_dir))

    workbook = openpyxl.load_workbook(excel_path, read_only=True, data_only=True)
    try:
        if sheet_name is None:
            worksheets = list(workbook.worksheets)
        else:
            if sheet_name not in workbook.sheetnames:
                raise ValueError('工作表不存在：{}，可选：{}'.format(sheet_name, workbook.sheetnames))
            worksheets = [workbook[sheet_name]]

        multi = len(worksheets) > 1
        stem = os.path.splitext(os.path.basename(excel_path))[0]
        files = []
        row_counts = {}

        for worksheet in worksheets:
            raw_rows = [list(row) for row in worksheet.iter_rows(values_only=True)]
            while raw_rows and all(cell is None or cell == '' for cell in raw_rows[-1]):
                raw_rows.pop()

            name = worksheet.title
            filename = '{}_{}.csv'.format(stem, _safe_name(name)) if multi else '{}.csv'.format(stem)
            file_path = os.path.join(target_dir, filename)

            with open(file_path, 'w', encoding=encoding, newline='') as stream:
                writer = csv.writer(stream, delimiter=delimiter)
                for row in raw_rows:
                    writer.writerow(['' if cell is None else cell for cell in row])

            files.append(os.path.abspath(file_path))
            row_counts[name] = max(len(raw_rows) - 1, 0)
    finally:
        workbook.close()

    return {
        'files': files,
        'rows': row_counts,
        'encoding': encoding,
        'delimiter': delimiter,
    }


def _safe_name(name):
    """把工作表名转换成安全的文件名片段。

    Args:
        name (str): 工作表名称。

    Returns:
        str: 替换掉非法字符并去掉首尾空格后的名称；为空时返回 'sheet'。
    """
    cleaned = _INVALID_FILENAME_CHARS.sub('_', str(name)).strip()
    return cleaned or 'sheet'
