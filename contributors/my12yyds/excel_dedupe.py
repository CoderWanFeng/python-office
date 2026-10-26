# -*- coding: UTF-8 -*-
"""按指定列给 Excel 去重，并给出重复明细。

适用场景：多份表格粘贴到一起后出现重复记录；或者导入系统前需要确认「哪些行是重复的」。
只给一个「去重后行数」往往不够，出问题时还要能指出重复的是哪几行、和哪一行重复了，
所以本方法同时返回重复明细，并可把去重结果写成新文件。

与 office.api.excel.merge2excel 的区别：merge2excel 只负责合并，不判断内容是否重复；
本方法专注于重复检测与去重。

Author:
    my12yyds

Project:
    https://www.python-office.com
"""

import os

import openpyxl


def excel_dedupe(
    excel_path: str,
    key_columns=None,
    sheet_name: str = None,
    keep: str = 'first',
    case_sensitive: bool = True,
    output_path: str = None,
) -> dict:
    """按指定列（或整行）给 Excel 去重。

    Args:
        excel_path (str): 需要去重的 Excel 文件路径，支持 .xlsx。
        key_columns (str | list, optional): 作为去重依据的列名，可以是单个列名字符串，
            也可以是列名列表。默认 None，表示用整行内容作为去重依据。
        sheet_name (str, optional): 工作表名称。默认 None，表示第一个工作表。
        keep (str, optional): 重复时保留哪一行，'first' 保留第一次出现的行，
            'last' 保留最后一次出现的行。默认 'first'。
        case_sensitive (bool, optional): 为 True 时文本比较区分大小写。
            默认 True；设为 False 时「ABC」与「abc」视为重复。
        output_path (str, optional): 去重结果的输出路径。指定后会写出一个只含保留行的
            .xlsx 文件，表头保持不变。默认 None，不写文件。

    Returns:
        dict: 去重结果，包含以下键：
            - sheet (str): 工作表名称。
            - key_columns (list): 实际使用的去重列；整行去重时为 None。
            - total (int): 数据行总数（不含表头）。
            - unique (int): 去重后保留的行数。
            - duplicate_count (int): 被判定为重复而剔除的行数。
            - duplicates (list): 重复明细，每项为 dict，含 row（重复行行号）、
              kept_row（保留行行号）、key（该行的去重键，列表）。
            - output_path (str): 去重结果文件路径；未导出时为 None。

    Raises:
        FileNotFoundError: Excel 文件不存在。
        ValueError: 工作表不存在、表格为空、keep 取值非法，或去重列不存在。
    """
    if not os.path.isfile(excel_path):
        raise FileNotFoundError('文件不存在：{}'.format(excel_path))
    if keep not in ('first', 'last'):
        raise ValueError("keep 只支持 'first' 或 'last'，收到：{}".format(keep))

    header, rows = _load_sheet(excel_path, sheet_name)
    if not header:
        raise ValueError('表格为空，至少需要一行表头：{}'.format(excel_path))

    if key_columns is None:
        positions = None
        key_names = None
    else:
        key_names = [key_columns] if isinstance(key_columns, str) else list(key_columns)
        missing = [name for name in key_names if name not in header]
        if missing:
            raise ValueError('去重列不存在：{}，可选列：{}'.format('、'.join(missing), '、'.join(header)))
        positions = [header.index(name) for name in key_names]

    groups = {}
    order = []
    for offset, row in enumerate(rows):
        raw_key = list(row) if positions is None else [row[i] if i < len(row) else None for i in positions]
        key = _normalize_key(raw_key, case_sensitive)
        if key not in groups:
            groups[key] = {'indices': [], 'display': _display_key(raw_key)}
            order.append(key)
        groups[key]['indices'].append(offset)

    kept_indices = []
    duplicates = []
    for key in order:
        item = groups[key]
        indices = item['indices']
        if keep == 'first':
            kept, dropped = indices[0], indices[1:]
        else:
            kept, dropped = indices[-1], indices[:-1]
        kept_indices.append(kept)
        for index in dropped:
            duplicates.append({
                'row': index + 2,
                'kept_row': kept + 2,
                'key': item['display'],
            })

    kept_indices.sort()
    duplicates.sort(key=lambda item: item['row'])

    result = {
        'sheet': _sheet_label(excel_path, sheet_name),
        'key_columns': key_names,
        'total': len(rows),
        'unique': len(kept_indices),
        'duplicate_count': len(duplicates),
        'duplicates': duplicates,
        'output_path': None,
    }

    if output_path:
        result['output_path'] = _write_unique(header, rows, kept_indices, output_path)
    return result


def _load_sheet(path, sheet_name):
    """读取工作表，返回（表头列表，数据行列表）。

    Args:
        path (str): Excel 文件路径。
        sheet_name (str): 工作表名称，None 表示第一个工作表。

    Returns:
        tuple: (header, rows)。rows 已去掉末尾的空行。

    Raises:
        ValueError: 指定的工作表不存在。
    """
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        if sheet_name is None:
            worksheet = workbook.worksheets[0]
        else:
            if sheet_name not in workbook.sheetnames:
                raise ValueError('工作表不存在：{}，可选：{}'.format(sheet_name, workbook.sheetnames))
            worksheet = workbook[sheet_name]
        raw_rows = [list(row) for row in worksheet.iter_rows(values_only=True)]
    finally:
        workbook.close()

    while raw_rows and all(cell is None or cell == '' for cell in raw_rows[-1]):
        raw_rows.pop()
    if not raw_rows:
        return [], []

    header = []
    for index, cell in enumerate(raw_rows[0]):
        text = str(cell).strip() if cell is not None else ''
        header.append(text or '列{}'.format(index + 1))
    return header, raw_rows[1:]


def _sheet_label(path, sheet_name):
    """返回用于结果展示的工作表名称。

    Args:
        path (str): Excel 文件路径。
        sheet_name (str): 用户指定的工作表名称。

    Returns:
        str: 工作表名称。
    """
    if sheet_name:
        return sheet_name
    workbook = openpyxl.load_workbook(path, read_only=True)
    try:
        return workbook.sheetnames[0]
    finally:
        workbook.close()


def _normalize_key(values, case_sensitive):
    """把一行值转换成可哈希、可比较的去重键。

    Args:
        values (list): 参与去重的单元格值。
        case_sensitive (bool): 为 True 时文本区分大小写。

    Returns:
        tuple: 去重键。空字符串与 None 视为同一个值。
    """
    key = []
    for value in values:
        if value is None or value == '':
            key.append(None)
        elif isinstance(value, str):
            text = value.strip()
            key.append(text if case_sensitive else text.lower())
        else:
            key.append(value)
    return tuple(key)


def _display_key(values):
    """返回用于展示的原始去重键。

    Args:
        values (list): 参与去重的单元格值。

    Returns:
        list: 原始值列表，None 转为空字符串。
    """
    return ['' if value is None else value for value in values]


def _write_unique(header, rows, kept_indices, output_path):
    """把保留的行写成新的 Excel 文件。

    Args:
        header (list): 表头列表。
        rows (list): 原始数据行。
        kept_indices (list): 需要保留的行下标，已按原顺序排序。
        output_path (str): 输出文件路径。

    Returns:
        str: 实际写入的文件路径。

    Raises:
        OSError: 输出目录不存在或不可写。
    """
    directory = os.path.dirname(os.path.abspath(output_path))
    if not os.path.isdir(directory):
        raise OSError('输出目录不存在：{}'.format(directory))

    workbook = openpyxl.Workbook()
    worksheet = workbook.active
    worksheet.title = '去重结果'
    worksheet.append(header)
    for index in kept_indices:
        worksheet.append(list(rows[index]))
    workbook.save(output_path)
    return os.path.abspath(output_path)
