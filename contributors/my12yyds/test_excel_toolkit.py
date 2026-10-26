# -*- coding: UTF-8 -*-
"""excel_diff / excel_dedupe / excel_column_stats / excel2json / excel2csv 的单元测试。

运行方式（在仓库根目录）：
    pytest contributors/my12yyds/test_excel_toolkit.py -v
"""

import datetime
import importlib.util
import json
import os
from pathlib import Path
import tempfile
import unittest

import openpyxl


def _load(module_name):
    """按文件路径加载同目录下的模块。

    Args:
        module_name (str): 不含 .py 的模块文件名。

    Returns:
        module: 已加载的模块对象。
    """
    module_path = Path(__file__).resolve().parent / '{}.py'.format(module_name)
    spec = importlib.util.spec_from_file_location('{}_under_test'.format(module_name), module_path)
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


excel_diff = _load('excel_diff').excel_diff
excel_dedupe = _load('excel_dedupe').excel_dedupe
excel_column_stats = _load('excel_column_stats').excel_column_stats
excel2json = _load('excel2json').excel2json
excel2csv = _load('excel2csv').excel2csv


def _make_excel(path, rows, header=None, sheet_name='Sheet1'):
    """生成一个用于测试的 Excel 文件。

    Args:
        path (str): 输出文件路径。
        rows (list): 二维数据行。
        header (list): 表头，None 表示不写表头。
        sheet_name (str): 工作表名称。

    Returns:
        str: 文件路径。
    """
    workbook = openpyxl.Workbook()
    worksheet = workbook.active
    worksheet.title = sheet_name
    if header:
        worksheet.append(list(header))
    for row in rows:
        worksheet.append(list(row))
    workbook.save(path)
    return path


class TestExcelDiff(unittest.TestCase):
    """excel_diff 的测试。"""

    def setUp(self):
        self.directory = tempfile.TemporaryDirectory()
        self.root = self.directory.name

    def tearDown(self):
        self.directory.cleanup()

    def _path(self, name):
        return os.path.join(self.root, name)

    def test_identical_tables(self):
        before = _make_excel(self._path('a.xlsx'), [['A', 1]], ['名称', '数量'])
        after = _make_excel(self._path('b.xlsx'), [['A', 1]], ['名称', '数量'])
        result = excel_diff(before, after)
        self.assertTrue(result['summary']['identical'])
        self.assertEqual(result['changes'], [])
        self.assertEqual(result['mode'], 'position')

    def test_position_mode_reports_changed_cell(self):
        before = _make_excel(self._path('a.xlsx'), [['A', 1], ['B', 2]], ['名称', '数量'])
        after = _make_excel(self._path('b.xlsx'), [['A', 1], ['B', 3]], ['名称', '数量'])
        result = excel_diff(before, after)
        self.assertFalse(result['summary']['identical'])
        self.assertEqual(result['summary']['changed'], 1)
        change = result['changes'][0]
        self.assertEqual(change['row'], 3)
        self.assertEqual(change['column'], '数量')
        self.assertEqual(change['before'], 2)
        self.assertEqual(change['after'], 3)

    def test_position_mode_reports_added_and_removed_rows(self):
        before = _make_excel(self._path('a.xlsx'), [['A'], ['B']], ['名称'])
        after = _make_excel(self._path('b.xlsx'), [['A'], ['C']], ['名称'])
        result = excel_diff(before, after)
        self.assertEqual(result['summary']['changed'], 1)
        self.assertEqual(result['added_rows'], [])
        self.assertEqual(result['removed_rows'], [])

        longer = _make_excel(self._path('c.xlsx'), [['A'], ['B'], ['C']], ['名称'])
        result = excel_diff(before, longer)
        self.assertEqual(len(result['added_rows']), 1)
        self.assertEqual(result['added_rows'][0]['row'], 4)

    def test_key_mode_matches_rows_by_key(self):
        before = _make_excel(
            self._path('a.xlsx'),
            [['A001', '张三', 100], ['A002', '李四', 200]],
            ['编号', '姓名', '金额'],
        )
        after = _make_excel(
            self._path('b.xlsx'),
            [['A002', '李四', 250], ['A003', '王五', 300]],
            ['编号', '姓名', '金额'],
        )
        result = excel_diff(before, after, key_column='编号')
        self.assertEqual(result['mode'], 'key')
        self.assertEqual([item['key'] for item in result['added_rows']], ['A003'])
        self.assertEqual([item['key'] for item in result['removed_rows']], ['A001'])
        self.assertEqual(len(result['changes']), 1)
        self.assertEqual(result['changes'][0]['key'], 'A002')
        self.assertEqual(result['changes'][0]['column'], '金额')

    def test_key_mode_ignores_row_order(self):
        before = _make_excel(self._path('a.xlsx'), [['A', 1], ['B', 2]], ['编号', '值'])
        after = _make_excel(self._path('b.xlsx'), [['B', 2], ['A', 1]], ['编号', '值'])
        result = excel_diff(before, after, key_column='编号')
        self.assertTrue(result['summary']['identical'])

    def test_key_mode_rejects_duplicate_keys(self):
        before = _make_excel(self._path('a.xlsx'), [['A', 1], ['A', 2]], ['编号', '值'])
        after = _make_excel(self._path('b.xlsx'), [['A', 1]], ['编号', '值'])
        with self.assertRaises(ValueError):
            excel_diff(before, after, key_column='编号')

    def test_key_mode_rejects_missing_key_column(self):
        before = _make_excel(self._path('a.xlsx'), [['A', 1]], ['编号', '值'])
        after = _make_excel(self._path('b.xlsx'), [['A', 1]], ['编号', '值'])
        with self.assertRaises(ValueError):
            excel_diff(before, after, key_column='不存在的列')

    def test_ignore_case(self):
        before = _make_excel(self._path('a.xlsx'), [['abc']], ['名称'])
        after = _make_excel(self._path('b.xlsx'), [['ABC']], ['名称'])
        self.assertFalse(excel_diff(before, after)['summary']['identical'])
        self.assertTrue(excel_diff(before, after, ignore_case=True)['summary']['identical'])

    def test_header_changes_are_reported(self):
        before = _make_excel(self._path('a.xlsx'), [['A', 1]], ['名称', '旧列'])
        after = _make_excel(self._path('b.xlsx'), [['A', 1]], ['名称', '新列'])
        result = excel_diff(before, after)
        self.assertEqual(result['header']['added'], ['新列'])
        self.assertEqual(result['header']['removed'], ['旧列'])

    def test_output_report_is_written(self):
        before = _make_excel(self._path('a.xlsx'), [['A', 1]], ['名称', '数量'])
        after = _make_excel(self._path('b.xlsx'), [['A', 2]], ['名称', '数量'])
        output = self._path('diff.xlsx')
        result = excel_diff(before, after, output_path=output)
        self.assertTrue(os.path.isfile(output))
        self.assertEqual(result['output_path'], os.path.abspath(output))
        workbook = openpyxl.load_workbook(output)
        self.assertEqual(workbook.sheetnames, ['差异', '汇总'])
        workbook.close()

    def test_missing_file(self):
        existing = _make_excel(self._path('a.xlsx'), [['A']], ['名称'])
        with self.assertRaises(FileNotFoundError):
            excel_diff(existing, self._path('nope.xlsx'))


class TestExcelDedupe(unittest.TestCase):
    """excel_dedupe 的测试。"""

    def setUp(self):
        self.directory = tempfile.TemporaryDirectory()
        self.root = self.directory.name

    def tearDown(self):
        self.directory.cleanup()

    def _path(self, name):
        return os.path.join(self.root, name)

    def test_whole_row_dedupe_keeps_first(self):
        path = _make_excel(self._path('a.xlsx'), [['A', 1], ['B', 2], ['A', 1]], ['名称', '数量'])
        result = excel_dedupe(path)
        self.assertEqual(result['total'], 3)
        self.assertEqual(result['unique'], 2)
        self.assertEqual(result['duplicate_count'], 1)
        self.assertEqual(result['duplicates'][0]['row'], 4)
        self.assertEqual(result['duplicates'][0]['kept_row'], 2)

    def test_key_column_as_string(self):
        path = _make_excel(self._path('a.xlsx'), [['A', 1], ['A', 2], ['B', 3]], ['编号', '值'])
        result = excel_dedupe(path, key_columns='编号')
        self.assertEqual(result['key_columns'], ['编号'])
        self.assertEqual(result['unique'], 2)
        self.assertEqual(result['duplicate_count'], 1)

    def test_multiple_key_columns(self):
        path = _make_excel(self._path('a.xlsx'), [['A', 1], ['A', 2], ['B', 2]], ['编号', '值'])
        result = excel_dedupe(path, key_columns=['编号', '值'])
        self.assertEqual(result['unique'], 3)

    def test_keep_last(self):
        path = _make_excel(self._path('a.xlsx'), [['A', 1], ['A', 2]], ['编号', '值'])
        result = excel_dedupe(path, key_columns='编号', keep='last')
        self.assertEqual(result['duplicates'][0]['row'], 2)
        self.assertEqual(result['duplicates'][0]['kept_row'], 3)

    def test_case_sensitivity(self):
        path = _make_excel(self._path('a.xlsx'), [['ABC'], ['abc']], ['名称'])
        self.assertEqual(excel_dedupe(path)['unique'], 2)
        self.assertEqual(excel_dedupe(path, case_sensitive=False)['unique'], 1)

    def test_blank_cells_are_treated_as_equal(self):
        path = _make_excel(self._path('a.xlsx'), [['A', None], ['A', '']], ['编号', '备注'])
        self.assertEqual(excel_dedupe(path)['unique'], 1)

    def test_output_file_contains_unique_rows(self):
        path = _make_excel(self._path('a.xlsx'), [['A', 1], ['B', 2], ['A', 1]], ['名称', '数量'])
        output = self._path('unique.xlsx')
        result = excel_dedupe(path, output_path=output)
        self.assertTrue(os.path.isfile(output))
        self.assertEqual(result['output_path'], os.path.abspath(output))
        workbook = openpyxl.load_workbook(output)
        rows = list(workbook.active.iter_rows(values_only=True))
        workbook.close()
        self.assertEqual(len(rows), 3)  # 表头 + 2 行
        self.assertEqual(rows[0], ('名称', '数量'))

    def test_invalid_arguments(self):
        path = _make_excel(self._path('a.xlsx'), [['A']], ['名称'])
        with self.assertRaises(ValueError):
            excel_dedupe(path, keep='middle')
        with self.assertRaises(ValueError):
            excel_dedupe(path, key_columns='不存在的列')
        with self.assertRaises(FileNotFoundError):
            excel_dedupe(self._path('nope.xlsx'))


class TestExcelColumnStats(unittest.TestCase):
    """excel_column_stats 的测试。"""

    def setUp(self):
        self.directory = tempfile.TemporaryDirectory()
        self.root = self.directory.name

    def tearDown(self):
        self.directory.cleanup()

    def test_numeric_column_statistics(self):
        path = _make_excel(
            os.path.join(self.root, 'a.xlsx'),
            [['A', 10], ['B', 20], ['C', 30]],
            ['名称', '数量'],
        )
        result = excel_column_stats(path)
        column = result['columns'][1]
        self.assertEqual(column['type'], 'number')
        self.assertEqual(column['numeric']['min'], 10)
        self.assertEqual(column['numeric']['max'], 30)
        self.assertEqual(column['numeric']['mean'], 20)
        self.assertEqual(column['numeric']['median'], 20)
        self.assertEqual(column['numeric']['stdev'], 10)
        self.assertEqual(result['row_count'], 3)
        self.assertEqual(result['fill_rate'], 1.0)

    def test_text_column_top_values(self):
        path = _make_excel(
            os.path.join(self.root, 'a.xlsx'),
            [['北京'], ['上海'], ['北京'], ['北京'], ['上海']],
            ['城市'],
        )
        result = excel_column_stats(path, top_n=2)
        column = result['columns'][0]
        self.assertEqual(column['type'], 'text')
        self.assertEqual(column['unique'], 2)
        self.assertEqual(column['top'][0], {'value': '北京', 'count': 3})
        self.assertEqual(column['top'][1], {'value': '上海', 'count': 2})

    def test_empty_column_and_empty_cells(self):
        path = _make_excel(
            os.path.join(self.root, 'a.xlsx'),
            [['A', None], ['B', None]],
            ['名称', '备注'],
        )
        result = excel_column_stats(path)
        remark = result['columns'][1]
        self.assertEqual(remark['type'], 'empty')
        self.assertEqual(remark['empty'], 2)
        self.assertEqual(remark['empty_ratio'], 1.0)
        self.assertEqual(result['fill_rate'], 0.5)

    def test_mixed_column(self):
        path = _make_excel(
            os.path.join(self.root, 'a.xlsx'),
            [[1], ['文本'], [3]],
            ['值'],
        )
        result = excel_column_stats(path)
        self.assertEqual(result['columns'][0]['type'], 'mixed')
        self.assertIn('top', result['columns'][0])

    def test_date_column(self):
        path = _make_excel(
            os.path.join(self.root, 'a.xlsx'),
            [[datetime.datetime(2024, 1, 1)], [datetime.datetime(2024, 3, 1)]],
            ['日期'],
        )
        result = excel_column_stats(path)
        column = result['columns'][0]
        self.assertEqual(column['type'], 'date')
        self.assertEqual(column['date_range']['min'], '2024-01-01T00:00:00')

    def test_top_n_zero_disables_top(self):
        path = _make_excel(os.path.join(self.root, 'a.xlsx'), [['A'], ['A']], ['名称'])
        result = excel_column_stats(path, top_n=0)
        self.assertNotIn('top', result['columns'][0])

    def test_negative_top_n(self):
        path = _make_excel(os.path.join(self.root, 'a.xlsx'), [['A']], ['名称'])
        with self.assertRaises(ValueError):
            excel_column_stats(path, top_n=-1)


class TestExcel2Json(unittest.TestCase):
    """excel2json 的测试。"""

    def setUp(self):
        self.directory = tempfile.TemporaryDirectory()
        self.root = self.directory.name

    def tearDown(self):
        self.directory.cleanup()

    def _sample(self):
        return _make_excel(
            os.path.join(self.root, 'a.xlsx'),
            [['A', 1], ['B', None]],
            ['名称', '数量'],
        )

    def test_records_orientation(self):
        data = excel2json(self._sample())
        self.assertEqual(data, [{'名称': 'A', '数量': 1}, {'名称': 'B', '数量': None}])

    def test_columns_orientation(self):
        data = excel2json(self._sample(), orient='columns')
        self.assertEqual(data['名称'], ['A', 'B'])
        self.assertEqual(data['数量'], [1, None])

    def test_values_orientation(self):
        data = excel2json(self._sample(), orient='values')
        self.assertEqual(data, [['A', 1], ['B', None]])

    def test_duplicate_headers_get_suffix(self):
        path = _make_excel(os.path.join(self.root, 'b.xlsx'), [['x', 'y']], ['值', '值'])
        data = excel2json(path)
        self.assertEqual(data, [{'值': 'x', '值_2': 'y'}])

    def test_datetime_becomes_iso_string(self):
        path = _make_excel(
            os.path.join(self.root, 'c.xlsx'),
            [[datetime.datetime(2024, 5, 6, 7, 8, 9)]],
            ['时间'],
        )
        self.assertEqual(excel2json(path)[0]['时间'], '2024-05-06T07:08:09')

    def test_output_file(self):
        output = os.path.join(self.root, 'out.json')
        data = excel2json(self._sample(), output_path=output, ensure_ascii=False)
        self.assertTrue(os.path.isfile(output))
        with open(output, encoding='utf-8') as stream:
            self.assertEqual(json.load(stream), data)

    def test_invalid_orient(self):
        with self.assertRaises(ValueError):
            excel2json(self._sample(), orient='table')


class TestExcel2Csv(unittest.TestCase):
    """excel2csv 的测试。"""

    def setUp(self):
        self.directory = tempfile.TemporaryDirectory()
        self.root = self.directory.name

    def tearDown(self):
        self.directory.cleanup()

    def test_single_sheet(self):
        path = _make_excel(os.path.join(self.root, 'a.xlsx'), [['张三', 1]], ['姓名', '数量'])
        result = excel2csv(path)
        self.assertEqual(len(result['files']), 1)
        self.assertTrue(result['files'][0].endswith('a.csv'))
        self.assertEqual(result['rows'], {'Sheet1': 1})

    def test_utf8_sig_bom_for_excel(self):
        path = _make_excel(os.path.join(self.root, 'a.xlsx'), [['张三']], ['姓名'])
        result = excel2csv(path)
        with open(result['files'][0], 'rb') as stream:
            self.assertTrue(stream.read().startswith(b'\xef\xbb\xbf'))

    def test_multiple_sheets(self):
        workbook = openpyxl.Workbook()
        first = workbook.active
        first.title = '一月'
        first.append(['姓名'])
        first.append(['张三'])
        second = workbook.create_sheet('二月')
        second.append(['姓名'])
        second.append(['李四'])
        path = os.path.join(self.root, 'multi.xlsx')
        workbook.save(path)

        result = excel2csv(path)
        self.assertEqual(len(result['files']), 2)
        self.assertTrue(any(name.endswith('multi_一月.csv') for name in result['files']))
        self.assertEqual(result['rows'], {'一月': 1, '二月': 1})

    def test_sheet_name_filter(self):
        workbook = openpyxl.Workbook()
        workbook.active.title = '一月'
        workbook.active.append(['姓名'])
        workbook.create_sheet('二月').append(['姓名'])
        path = os.path.join(self.root, 'multi.xlsx')
        workbook.save(path)

        result = excel2csv(path, sheet_name='二月')
        self.assertEqual(len(result['files']), 1)
        self.assertTrue(result['files'][0].endswith('multi.csv'))

    def test_custom_delimiter_and_encoding(self):
        path = _make_excel(os.path.join(self.root, 'a.xlsx'), [['a;b', 'c']], ['一', '二'])
        result = excel2csv(path, encoding='utf-8', delimiter=';')
        with open(result['files'][0], encoding='utf-8') as stream:
            content = stream.read()
        self.assertIn('"a;b";c', content)

    def test_missing_output_directory(self):
        path = _make_excel(os.path.join(self.root, 'a.xlsx'), [['A']], ['名称'])
        with self.assertRaises(FileNotFoundError):
            excel2csv(path, output_dir=os.path.join(self.root, 'nope'))


if __name__ == '__main__':
    unittest.main()
