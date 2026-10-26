# -*- coding: UTF-8 -*-
"""把 Excel 工作表转换成 JSON。

适用场景：把运营同学维护的 Excel 交给程序消费——生成接口测试数据、导入数据库、
或者塞进前端页面的 mock 数据。手工另存为 CSV 再转 JSON 很容易在日期、空值和中文编码上踩坑。

三种输出结构：
    records —— 每行一个对象，最常用，适合直接当接口数据。
    columns —— 每列一个数组，适合按字段批量处理。
    values  —— 纯二维数组，不含表头，适合喂给表格组件。

与 office.api.markdown.excel2markdown 的区别：excel2markdown 输出给人看的 Markdown 表格，
本方法输出给程序读的 JSON。

Author:
    my12yyds

Project:
    https://www.python-office.com
"""

import datetime
import json
import os

import openpyxl

#: 支持的输出结构。
ORIENTS = ('records', 'columns', 'values')


def excel2json(
    excel_path: str,
    sheet_name: str = None,
    orient: str = 'records',
    output_path: str = None,
    indent: int = 2,
    ensure_ascii: bool = False,
):
    """把工作表转换成 JSON 兼容的 Python 对象，可选写入文件。

    Args:
        excel_path (str): Excel 文件路径，支持 .xlsx。
        sheet_name (str, optional): 工作表名称。默认 None，表示第一个工作表。
        orient (str, optional): 输出结构，可选 'records'、'columns'、'values'。
            默认 'records'。
        output_path (str, optional): JSON 文件输出路径。指定后会同时写入文件。
            默认 None，只返回对象不写文件。
        indent (int, optional): 写入文件时的缩进空格数，设为 None 表示压缩成一行。
            默认 2。
        ensure_ascii (bool, optional): 为 True 时中文写成 \\uXXXX 转义。
            默认 False，保留中文原文，文件更易读。

    Returns:
        list | dict: 转换结果。
            - orient='records' 时为 list，每项是「列名 -> 值」的 dict。
            - orient='columns' 时为 dict，每项是「列名 -> 值列表」。
            - orient='values' 时为 list，每项是一行的值列表，不含表头。

    Raises:
        FileNotFoundError: Excel 文件不存在。
        ValueError: 工作表不存在、表格为空，或 orient 取值非法。
        OSError: 指定了 output_path 但输出目录不存在或不可写。
    """
    if not os.path.isfile(excel_path):
        raise FileNotFoundError('文件不存在：{}'.format(excel_path))
    if orient not in ORIENTS:
        raise ValueError('orient 只支持 {}，收到：{}'.format('、'.join(ORIENTS), orient))

    header, rows = _load_sheet(excel_path, sheet_name)
    if not header:
        raise ValueError('表格为空，至少需要一行表头：{}'.format(excel_path))

    if orient == 'records':
        data = [
            {name: _json_value(row[index] if index < len(row) else None) for index, name in enumerate(header)}
            for row in rows
        ]
    elif orient == 'columns':
        data = {
            name: [_json_value(row[index] if index < len(row) else None) for row in rows]
            for index, name in enumerate(header)
        }
    else:
        width = len(header)
        data = [
            [_json_value(row[index] if index < len(row) else None) for index in range(width)]
            for row in rows
        ]

    if output_path:
        directory = os.path.dirname(os.path.abspath(output_path))
        if not os.path.isdir(directory):
            raise OSError('输出目录不存在：{}'.format(directory))
        with open(output_path, 'w', encoding='utf-8') as stream:
            json.dump(data, stream, indent=indent, ensure_ascii=ensure_ascii)
            stream.write('\n')

    return data


def _load_sheet(path, sheet_name):
    """读取工作表，返回（表头列表，数据行列表）。

    Args:
        path (str): Excel 文件路径。
        sheet_name (str): 工作表名称，None 表示第一个工作表。

    Returns:
        tuple: (header, rows)。表头中的重复列名会自动加 _2、_3 后缀，
        避免转换成 dict 时互相覆盖。

    Raises:
        ValueError: 指定的工作表不存在。
    """
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        if sheet_name is None:
            worksheet = workbook.worksheets[0]
        else:
            if sheet_name not in workbook.sheetnames:
                raise ValueError('工作表不存在：{}，可选：{}'.format(sheet_name, workbook.sheetnames))
            worksheet = workbook[sheet_name]
        raw_rows = [list(row) for row in worksheet.iter_rows(values_only=True)]
    finally:
        workbook.close()

    while raw_rows and all(cell is None or cell == '' for cell in raw_rows[-1]):
        raw_rows.pop()
    if not raw_rows:
        return [], []

    header = []
    seen = {}
    for index, cell in enumerate(raw_rows[0]):
        text = str(cell).strip() if cell is not None else ''
        text = text or '列{}'.format(index + 1)
        if text in seen:
            seen[text] += 1
            text = '{}_{}'.format(text, seen[text])
        else:
            seen[text] = 1
        header.append(text)
    return header, raw_rows[1:]


def _json_value(value):
    """把单元格值转换成 JSON 可序列化的值。

    Args:
        value (object): 任意单元格值。

    Returns:
        object: 空单元格返回 None；日期时间返回 ISO 字符串；
        数字、布尔、字符串原样返回；其他类型转为字符串。
    """
    if value is None:
        return None
    if isinstance(value, str):
        return value
    if isinstance(value, bool):
        return value
    if isinstance(value, (int, float)):
        return value
    if isinstance(value, (datetime.datetime, datetime.date, datetime.time)):
        return value.isoformat()
    return str(value)
