"""Tests for sheet name handling in the bundled Excel splitter."""

import importlib.util
from pathlib import Path
import tempfile
import unittest

import openpyxl


def load_split_excel_module():
    module_path = Path(__file__).resolve().parents[2] / "office" / "lib" / "excel" / "SplitExcel.py"
    spec = importlib.util.spec_from_file_location("split_excel_under_test", module_path)
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


class TestSplitExcelSheetName(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.split_excel = load_split_excel_module()

    def test_generate_xlsx_does_not_leave_empty_default_sheet(self):
        """输出文件不应残留 openpyxl 自动创建的空白 Sheet。"""
        worksheet_data = {
            "北京": [["张三", "北京", 100]],
            "上海": [["李四", "上海", 200]],
        }

        with tempfile.TemporaryDirectory() as tmp:
            filepath = str(Path(tmp) / "demo.xlsx")
            output = self.split_excel.generate_xlsx(filepath, worksheet_data)
            workbook = openpyxl.load_workbook(output)
            sheetnames = workbook.sheetnames

        self.assertEqual(sheetnames, ["北京", "上海"])

    def test_generate_xlsx_replaces_invalid_sheet_name_characters(self):
        """分组值包含 Excel 不允许的字符时不应抛异常。"""
        worksheet_data = {"2026/10/22": [["张三", "2026/10/22"]]}

        with tempfile.TemporaryDirectory() as tmp:
            filepath = str(Path(tmp) / "demo.xlsx")
            output = self.split_excel.generate_xlsx(filepath, worksheet_data)
            sheetnames = openpyxl.load_workbook(output).sheetnames

        self.assertEqual(sheetnames, ["2026_10_22"])

    def test_generate_xls_accepts_long_sheet_name(self):
        """分组值超过 31 个字符时，.xls 分支不应抛 invalid worksheet name。"""
        worksheet_data = {"A" * 40: [[1, "A" * 40]]}

        with tempfile.TemporaryDirectory() as tmp:
            filepath = str(Path(tmp) / "demo.xls")
            output = self.split_excel.generate_xls(filepath, worksheet_data)
            exists = Path(output).exists()

        self.assertTrue(exists)

    def test_invalid_characters_are_replaced(self):
        used_names = set()

        self.assertEqual(
            self.split_excel._to_valid_sheet_name("a:b*c?d[e]f\\g", used_names),
            "a_b_c_d_e_f_g",
        )

    def test_long_name_is_truncated_to_31_chars(self):
        used_names = set()

        name = self.split_excel._to_valid_sheet_name("A" * 40, used_names)

        self.assertEqual(len(name), 31)

    def test_names_stay_unique_after_truncation(self):
        used_names = set()

        first = self.split_excel._to_valid_sheet_name("A" * 40 + "1", used_names)
        second = self.split_excel._to_valid_sheet_name("A" * 40 + "2", used_names)

        self.assertNotEqual(first, second)
        self.assertLessEqual(len(second), 31)

    def test_blank_value_falls_back_to_default_name(self):
        used_names = set()

        self.assertEqual(self.split_excel._to_valid_sheet_name("   ", used_names), "Sheet")


if __name__ == "__main__":
    unittest.main()
